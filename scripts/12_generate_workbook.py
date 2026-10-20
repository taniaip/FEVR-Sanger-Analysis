#!/usr/bin/env python3

"""Combine Sanger analysis CSV outputs into a formatted Excel workbook."""

import argparse
import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


SHEETS = [
    ("Candidates", "candidates/{s}_Preliminary_Candidates.csv"),
    ("Candidate_Summary", "candidates/{s}_Candidate_Evidence_Summary.csv"),
    ("SNV_Evidence", "variants/{s}_SNV_Evidence_Audit.csv"),
    ("Indel_Quality", "variants/{s}_Indel_Quality_Evidence_Audit.csv"),
    ("Raw_SNVs", "variants/{s}_Preliminary_SNV_Read_Evidence.csv"),
    ("Raw_Indels", "variants/{s}_Preliminary_Indel_Read_Evidence.csv"),
    ("Target_Coverage", "coverage/{s}_KIF11_E10_Position_Coverage.csv"),
    ("Alignment_Details", "coverage/{s}_KIF11_E10_Alignment_Details.csv"),
    ("Alignment_QC", "alignment/{s}_Read_Alignment_QC.csv"),
    ("Duplicate_Audit", "duplicates/{s}_AB1_Duplicate_Audit.csv"),
    ("Plasmid_Primer_QC", "plasmid/{s}_MAP3K6_Plasmid_Primer_QC.csv"),
    ("Trim_QC", "trimmed/{s}_Quality_Trim_Summary.csv"),
]


def load_csv(path):
    with path.open(newline="", encoding="utf-8-sig") as handle:
        rows = list(csv.reader(handle))

    if not rows:
        raise ValueError(f"Empty CSV: {path}")

    return rows


def convert(value):
    if value == "":
        return None
    if value == "True":
        return True
    if value == "False":
        return False
    return value


def style_header(sheet, row_number=1):
    for cell in sheet[row_number]:
        cell.fill = PatternFill(
            fill_type="solid",
            fgColor="17324D"
        )
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(
            vertical="center",
            wrap_text=True
        )

    sheet.row_dimensions[row_number].height = 30


def add_sheet(workbook, name, rows):
    sheet = workbook.create_sheet(name)

    for row in rows:
        sheet.append([convert(value) for value in row])

    style_header(sheet)

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions

    for column in sheet.columns:
        letter = get_column_letter(column[0].column)
        header = str(column[0].value or "")

        sheet.column_dimensions[letter].width = min(
            max(len(header) + 3, 16),
            38
        )

    count = len(rows) - 1
    print(f"  {name}: {count} rows")

    return count


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--results-dir", required=True, type=Path)
    parser.add_argument("--sample", required=True)
    parser.add_argument("--output", required=True, type=Path)

    args = parser.parse_args()

    results_dir = args.results_dir.expanduser().resolve()
    output = args.output.expanduser().resolve()

    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"

    summary_rows = [
        ["FEVR Sanger Sequencing — Analysis Summary", ""],
        ["Sample", args.sample],
        ["Reference", "GRCh38 chr10 (KIF11 genomic reads)"],
        ["Analysis status", "PRELIMINARY / UNVALIDATED"],
        ["Interpretation", "Read-level evidence; not clinical variant calls"],
        ["Plasmid analysis", "Primer QC only; no plasmid variant calling"],
        ["Indel representation", "Left-anchored, not normalized"],
        ["Biological independence", "Not established by distinct AB1 hashes"],
        ["", ""],
        ["Metric", "Count"],
    ]

    row_counts = {}

    for name, template in SHEETS:
        path = results_dir / template.format(s=args.sample)

        if not path.is_file():
            parser.error(f"Missing required CSV: {path}")

        rows = load_csv(path)
        row_counts[name] = add_sheet(workbook, name, rows)

    summary_rows.extend([
        ["Raw SNV observations", row_counts["Raw_SNVs"]],
        ["Distinct SNV alleles", row_counts["SNV_Evidence"]],
        ["Raw indel observations", row_counts["Raw_Indels"]],
        ["Candidate observation rows", row_counts["Candidates"]],
        ["Covered genomic positions", row_counts["Target_Coverage"]],
        ["Passing alignment detail rows", row_counts["Alignment_Details"]],
        ["AB1 files audited", row_counts["Duplicate_Audit"]],
        ["Plasmid primer groups", row_counts["Plasmid_Primer_QC"]],
    ])

    for row in summary_rows:
        summary.append(row)

    style_header(summary, 1)

    for cell in summary[10]:
        cell.fill = PatternFill(
            fill_type="solid",
            fgColor="DDEAF3"
        )
        cell.font = Font(color="17324D", bold=True)

    summary.column_dimensions["A"].width = 44
    summary.column_dimensions["B"].width = 55

    summary.freeze_panes = "A2"

    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)

    print("\nCONSOLIDATED WORKBOOK")
    print("=" * 65)
    print("Sample:", args.sample)
    print("Worksheets:", len(workbook.sheetnames))
    print("Output:", output)
    print("\nAll preliminary variants remain UNVALIDATED.")


if __name__ == "__main__":
    main()
